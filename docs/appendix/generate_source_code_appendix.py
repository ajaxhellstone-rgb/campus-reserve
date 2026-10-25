from __future__ import annotations

import subprocess
from collections import defaultdict
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = Path(__file__).resolve().parent / "Campus_Reserve_Source_Code_Appendix.docx"
SOURCE_ROOTS = (
    "app/",
    "bootstrap/",
    "config/",
    "database/",
    "docker/",
    "public/",
    "resources/",
    "routes/",
    "tests/",
)
ROOT_FILES = {
    ".env.example",
    "Dockerfile",
    "artisan",
    "composer.json",
    "package.json",
    "phpunit.xml",
    "vite.config.js",
}
SOURCE_EXTENSIONS = {
    ".conf",
    ".css",
    ".example",
    ".htaccess",
    ".js",
    ".json",
    ".php",
    ".sh",
    ".xml",
    ".yaml",
    ".yml",
}
EXCLUDED_PARTS = {"build", "cache", "node_modules", "output", "storage", "vendor"}


def tracked_source_files() -> list[Path]:
    tracked = subprocess.check_output(
        ["git", "ls-files", "-z"],
        cwd=ROOT,
    ).decode("utf-8").split("\0")
    selected: list[Path] = []

    for name in tracked:
        if not name or name in {".env", ".env.local"}:
            continue
        if not (name.startswith(SOURCE_ROOTS) or name in ROOT_FILES):
            continue

        path = Path(name)
        if path.suffix.lower() not in SOURCE_EXTENSIONS and name not in ROOT_FILES:
            continue
        if EXCLUDED_PARTS.intersection(path.parts):
            continue
        if path.is_file():
            selected.append(path)

    return sorted(selected, key=lambda path: path.as_posix().casefold())


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Page ")
    run.font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def format_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9)

    for style_name, size in (("Heading 1", 16), ("Heading 2", 10)):
        style = document.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(20, 83, 45)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = "CAMPUS RESERVE  |  SOURCE CODE APPENDIX"
    header.style = document.styles["Caption"]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number(section.footer.paragraphs[0])


def add_code_listing(document: Document, relative_path: Path) -> None:
    heading = document.add_paragraph(style="Heading 2")
    heading.paragraph_format.space_before = Pt(8)
    heading.paragraph_format.space_after = Pt(2)
    heading.add_run(relative_path.as_posix())

    content = (ROOT / relative_path).read_text(encoding="utf-8", errors="replace")
    lines = content.splitlines()
    if not lines:
        lines = [""]

    for number, line in enumerate(lines, start=1):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.line_spacing = 1
        paragraph.paragraph_format.keep_together = False

        run = paragraph.add_run(f"{number:>4}  {line}")
        run.font.name = "Consolas"
        run.font.size = Pt(7)
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")


def main() -> None:
    files = tracked_source_files()
    if not files:
        raise RuntimeError("No tracked source files were selected for the appendix.")

    document = Document()
    format_document(document)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(100)
    title_run = title.add_run("APPENDIX A")
    title_run.bold = True
    title_run.font.size = Pt(18)
    title_run.font.color.rgb = RGBColor(20, 83, 45)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("Campus Reserve Source Code Listing")
    subtitle_run.bold = True
    subtitle_run.font.size = Pt(24)

    description = document.add_paragraph()
    description.alignment = WD_ALIGN_PARAGRAPH.CENTER
    description.paragraph_format.space_before = Pt(12)
    description.add_run(
        f"Generated from {len(files)} tracked, first-party source and configuration files."
    )

    commit = subprocess.check_output(
        ["git", "rev-parse", "--short", "HEAD"],
        cwd=ROOT,
        text=True,
    ).strip()
    generated = document.add_paragraph()
    generated.alignment = WD_ALIGN_PARAGRAPH.CENTER
    generated.paragraph_format.space_before = Pt(6)
    generated.add_run(f"Repository revision: {commit}")
    document.add_page_break()

    document.add_heading("Scope and exclusions", level=1)
    document.add_paragraph(
        "This appendix provides the application source code, database migrations, "
        "views and frontend code, routes, tests, and selected deployment/configuration "
        "files for review and documentation."
    )
    document.add_paragraph(
        "Private environment files (including .env), credentials, dependencies "
        "(vendor and node_modules), runtime data, generated build output, binary "
        "assets, and unrelated documents are intentionally excluded. The tracked "
        ".env.example contains configuration placeholders only."
    )

    document.add_heading("File inventory", level=1)
    grouped: dict[str, list[str]] = defaultdict(list)
    for path in files:
        normalized = path.as_posix()
        group = normalized.split("/", 1)[0] if "/" in normalized else "Project root"
        grouped[group].append(normalized)

    inventory = document.add_table(rows=0, cols=2)
    inventory.style = "Light Shading Accent 1"
    for group in sorted(grouped, key=str.casefold):
        row = inventory.add_row()
        row.cells[0].text = group
        row.cells[1].text = f"{len(grouped[group])} files"
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, "E8F2EB")
            for run in cell.paragraphs[0].runs:
                run.bold = True
        for path in grouped[group]:
            row = inventory.add_row()
            row.cells[0].text = path
            row.cells[1].text = ""
            row.cells[0].paragraphs[0].paragraph_format.space_after = Pt(0)

    document.add_page_break()
    document.add_heading("Source code listings", level=1)
    document.add_paragraph(
        "Listings are reproduced from the tracked repository files. Line numbers "
        "are included for reference and are not part of the source code."
    )
    for path in files:
        add_code_listing(document, path)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(f"Created {OUTPUT}")
    print(f"Included {len(files)} files from revision {commit}")


if __name__ == "__main__":
    main()
